"""Build the comprehensive report and a self contained figure gallery."""
from pathlib import Path
import sys
import json
import html
import base64

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs"
OUT.mkdir(parents=True,exist_ok=True)
deps=ROOT/".build"/"python_deps"
if not deps.exists():
    deps=ROOT/"outputs"/"purchase_price_research"/".build"/"python_deps"
if deps.exists(): sys.path.insert(0,str(deps))
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

S=json.loads((ROOT/"tables"/"analysis_summary.json").read_text())
A=S["audit"]
D=S["diagnostics"]
M={r["model"]:r for r in S["metrics"]}
F={r["id"]:r for r in S["figures"]}
LABELS={"trade_count":"Trade count","unique_markets":"Markets traded","wallet_age_days":"Observed activity span","concentration":"Spending concentration","total_notional":"Total spending","max_single_buy":"Largest single purchase"}
PRIMARY=["Mean baseline","Linear regression","Gradient boosting"]
REFERENCE=[
    ("R1","Wallet feature table","Local course dataset at wallet_analysis/data/wallet_features.csv. SHA256 "+A["source_sha256"]+"."),
    ("R2","Course data definitions","L3 prediction markets student package, notebooks/insider_demo.ipynb, opening data description and feature dictionary. Used for cohort scope and variable definitions only. The insider detection task and its labels are outside this study."),
    ("R3","Gradient boosting implementation","https://scikit-learn.org/1.7/modules/generated/sklearn.ensemble.HistGradientBoostingRegressor.html"),
    ("R4","Cross validation","https://scikit-learn.org/1.7/modules/cross_validation.html"),
    ("R5","Permutation importance","https://scikit-learn.org/1.7/modules/permutation_importance.html"),
    ("R6","Preprocessing and leakage","https://scikit-learn.org/1.7/common_pitfalls.html"),
]

DOC=Document()
sec=DOC.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=DOC.styles
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Caption"]:
    styles[name].font.name="Arial"
    styles[name].font.color.rgb=RGBColor(0,0,0)
    styles[name]._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),"Arial")
    styles[name]._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),"Arial")
    for border in styles[name]._element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
styles["Normal"].font.size=Pt(11)
styles["Normal"].paragraph_format.line_spacing=1.08
styles["Normal"].paragraph_format.space_after=Pt(7)
styles["Normal"].paragraph_format.widow_control=True
styles["Title"].font.size=Pt(23)
styles["Title"].paragraph_format.space_after=Pt(9)
styles["Heading 1"].font.size=Pt(17)
styles["Heading 1"].paragraph_format.space_before=Pt(0)
styles["Heading 1"].paragraph_format.space_after=Pt(9)
styles["Heading 2"].font.size=Pt(12)
styles["Heading 2"].paragraph_format.space_before=Pt(8)
styles["Heading 2"].paragraph_format.space_after=Pt(5)
styles["Caption"].font.size=Pt(9)
styles["Caption"].font.italic=False
styles["Caption"].font.bold=False
styles["Caption"].paragraph_format.space_after=Pt(8)
header=sec.header.paragraphs[0]
header.text="Wallet purchase price analysis"
header.style=styles["Normal"]
header.runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run("STATS C161    ")
r.font.size=Pt(8)
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
foot._p.append(field)
DOC.core_properties.title="Wallet purchase price analysis"
DOC.core_properties.subject="Linear regression and gradient boosting on wallet trading characteristics"
DOC.core_properties.author=""
DOC.core_properties.keywords="wallets, regression, prediction, purchase price"
HTML=[]
MD=[]


def section(title,first=False):
    if not first: DOC.add_page_break()
    DOC.add_paragraph(title,style="Title" if first else "Heading 1")
    HTML.append(f"<section><h{'1' if first else '2'}>{html.escape(title)}</h{'1' if first else '2'}>")
    MD.append(("# " if first else "\n## ")+title+"\n")


def sub(title):
    DOC.add_paragraph(title,style="Heading 2")
    HTML.append("<h3>"+html.escape(title)+"</h3>")
    MD.append("\n### "+title+"\n")


def para(text,boldlead=None):
    p=DOC.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True;p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    HTML.append("<p>"+html.escape(text)+"</p>")
    MD.append(text+"\n")
    return p


def tab(headers,rows,widths=None):
    t=DOC.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None: widths=[7.1/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for cell,text,width in zip(t.rows[0].cells,headers,widths):
        cell.text=str(text);cell.width=Inches(width)
    for row in rows:
        cells=t.add_row().cells
        for cell,text,width in zip(cells,row,widths):
            cell.text=str(text);cell.width=Inches(width)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        no_split=OxmlElement("w:cantSplit");pr.append(no_split)
        if i==0:
            repeat=OxmlElement("w:tblHeader");pr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"263F5C" if i==0 else ("F0F4F8" if i%2==0 else "FFFFFF"));tcPr.append(shd)
            margins=OxmlElement("w:tcMar")
            for edge in ["top","left","bottom","right"]:
                node=OxmlElement("w:"+edge);node.set(qn("w:w"),"90");node.set(qn("w:type"),"dxa");margins.append(node)
            tcPr.append(margins)
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                node=OxmlElement("w:"+edge);node.set(qn("w:val"),"single");node.set(qn("w:sz"),"4");node.set(qn("w:color"),"D9D9D9");borders.append(node)
            tcPr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (i>0 and len(cell.text)>55) else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name="Arial";r.font.size=Pt(9.5);r.font.bold=i==0
                    r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    DOC.add_paragraph().paragraph_format.space_after=Pt(1)
    HTML.append("<table><thead><tr>"+"".join("<th>"+html.escape(str(h))+"</th>" for h in headers)+"</tr></thead><tbody>"+"".join("<tr>"+"".join("<td>"+html.escape(str(c))+"</td>" for c in row)+"</tr>" for row in rows)+"</tbody></table>")
    MD.append("| "+" | ".join(headers)+" |\n| "+" | ".join(["---"]*len(headers))+" |\n"+"\n".join("| "+" | ".join(str(c) for c in row)+" |" for row in rows)+"\n")


def fig(stem,width=6.9,max_height=None,caption=None):
    fp=ROOT/"figures"/(stem+".png")
    with Image.open(fp) as im: w,h=im.size
    if max_height and width*h/w>max_height: width=max_height*w/h
    p=DOC.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
    run=p.add_run();run.add_picture(str(fp),width=Inches(width))
    run._r.xpath(".//wp:docPr")[0].set("descr",F[stem]["title"])
    text=caption or F[stem]["caption"]
    DOC.add_paragraph("Figure "+stem[:2]+". "+text,style="Caption")
    b64=base64.b64encode(fp.read_bytes()).decode()
    HTML.append(f'<figure><img src="data:image/png;base64,{b64}" alt="{html.escape(F[stem]["title"])}"><figcaption>Figure {stem[:2]}. {html.escape(text)}</figcaption></figure>')
    MD.append("!["+F[stem]["title"]+"](../figures/"+stem+".png)\n\nFigure "+stem[:2]+". "+text+"\n")


section("Wallet purchase price analysis",True)
para("Research question: How well do wallet trading characteristics predict average purchase price, and does gradient boosting improve prediction over linear regression?")
para(f"Gradient boosting predicts average purchase price more accurately than the linear model in this dataset. On {A['test_rows']:,} held out wallets, its mean absolute error is {M['Gradient boosting']['MAE']*100:.2f} cents per share, compared with {M['Linear regression']['MAE']*100:.2f} cents for linear regression and {M['Mean baseline']['MAE']*100:.2f} cents for a constant mean prediction. The improvement over linear regression is 15.4%.")
tab(["Model","Test MAE cents","Test RMSE cents","Test R squared"],[[n,f"{M[n]['MAE']*100:.2f}",f"{M[n]['RMSE']*100:.2f}",f"{M[n]['R2']:.4f}"] for n in PRIMARY],[2.25,1.6,1.7,1.55])
fig("06_test_model_comparison",max_height=3.45)
para("We retain all 55,952 wallets and use six predictors. The confirmed flag plays no role in data loading, modeling, evaluation, or interpretation. The analysis describes prediction across wallets using summaries from the same trading period. It does not establish future forecasting ability, causal effects, profitability, or trading skill.")
para("Prepared October 26 2026. All reported model results come from the supplied feature table. The accompanying figure gallery provides all 15 figures as separate PNG and SVG files.")

section("Data scope and variable definitions")
para("The course notebook describes a cohort of wallets that bought shares in five selected event groups between April 2025 and March 2026: US strikes on Iran, a US and Iran ceasefire, the Machado Nobel Peace Prize, Maduro leaving office, and the OpenAI browser. Wallet features summarize their observed history in the lab trade files. These wallets form a selected cohort rather than a random sample of every platform user. [R1, R2]")
tab(["Field","Meaning","Unit"],[
    ["avg_buy_price","Volume weighted average price paid for purchased shares. This is the prediction target.","Dollars per share"],
    ["trade_count","Recorded trading activity count.","Count"],
    ["unique_markets","Number of distinct markets traded.","Count"],
    ["wallet_age_days","Time between the first and last observed trade. This measures activity span rather than account creation age.","Days"],
    ["concentration","Largest share of the wallet spending allocated to one market. A value of one means all spending goes to one market.","Proportion"],
    ["total_notional","Total recorded spending.","Dollars"],
    ["max_single_buy","Dollar value of the largest single purchase.","Dollars"],
],[1.55,4.4,1.15])
para("Each row represents one pseudonymous wallet. Wallet IDs serve only to check uniqueness and preserve split membership. Each wallet receives equal weight in model fitting and evaluation. The target uses volume weights within a wallet, while the study metrics use equal weights across wallets.")
tab(["Data check","Result"],[["Rows and unique wallets","55,952 each"],["Missing or nonfinite selected numeric values","0"],["Duplicate wallet IDs","0"],["Rows excluded","0"],["Wallets with one market","3,834"],["Wallets with zero observed span","1,184"]],[4.8,2.3])
para("The source has 11 columns. We load only the wallet ID, target, and six predictors. We exclude market_notional_ratio to keep the agreed specification focused. We exclude position_anomaly_ratio because its definition reuses largest purchase, spending, and market count. Neither exclusion depends on test performance.")

section("Data checks and preparation")
para("We check unique wallet IDs, complete finite numeric values, target prices between zero and one, and nonnegative inputs. We also check that market count does not exceed trade count, the largest purchase does not exceed total spending, and concentration is at least the inverse of the market count within rounding tolerance. All checks pass. No imputation, deduplication, winsorization, or outlier removal is necessary.")
para("We apply the natural logarithm of one plus the original value to trade count, market count, observed span, total spending, and largest purchase. This keeps zero spans valid and reduces the influence of long right tails. Concentration stays on its original scale. We leave the target on its original dollar scale.")
fig("02_predictor_distributions",max_height=4.0)
para("The linear model then standardizes each transformed predictor using means and standard deviations learned from its training fold. The boosting model uses the same transformed inputs without standardization. Log transformations mainly help the linear specification. Every fitted preprocessing step stays inside the model pipeline. [R6]")
para("All wallets remain in the study, including those that satisfy the demo activity threshold. That threshold defines a supplementary error subgroup only. We make no claim that high activity proves automated trading.")

section("Model design and evaluation")
sub("Training and test design")
para(f"We randomly assign 80% of wallets to training and 20% to a final test set, using seed 42. This yields {A['train_rows']:,} training wallets and {A['test_rows']:,} test wallets with no ID overlap. The same split applies to every model. Within the training set, five shuffled folds with seed 42 provide validation errors. All model settings and supplementary specifications are fixed before test evaluation. [R4]")
sub("Three primary comparisons")
para(f"The mean baseline assigns every test wallet the training target mean, {A['train_target_mean']:.6f} dollars per share. Linear regression fits an intercept and six additive terms using ordinary least squares on the transformed and standardized predictors. Histogram gradient boosting adds small regression trees sequentially, allowing nonlinear patterns and interactions. Both learned models fit squared error. We choose the boosting configuration using validation MAE. [R3]")
tab(["Boosting setting","Rule"],[["Learning rate","0.06"],["Boosting iterations","250"],["Minimum wallets per leaf","30"],["Maximum histogram bins","255"],["Maximum leaves searched","7, 15, 31"],["L2 regularization searched","0, 1"],["Early stopping","Disabled"],["Selection criterion","Lowest mean validation MAE across five folds"]],[3.3,3.8])
sub("Performance measures")
para("MAE is the average absolute difference between observed and predicted price and is the primary measure. RMSE gives extra weight to large errors. Multiplying either by 100 expresses the error in cents per share. R squared compares squared errors with the variation around the test target mean. A small negative baseline R squared can occur because its prediction uses the training mean rather than the test mean.")
para("We report primary predictions without clipping. We quantify MAE uncertainty with 2,000 paired bootstrap resamples of test wallets using seed 2026. Each resample preserves the pairing of errors across models. These intervals condition on the fitted models and assume wallet independence. They do not include training uncertainty or dependence among wallets sharing markets.")

section("Distribution of the prediction target")
para(f"The target spans {A['target_min']:.3f} to {A['target_max']:.3f} dollars per share. In training, the mean is {A['train_target_mean']:.3f} and the median is {A['train_target_median']:.3f}. The distribution has substantial mass toward higher prices and a smaller group at low prices. A constant mean therefore misses meaningful differences between wallets.")
fig("01_target_distribution",max_height=4.2)
split=pd.read_csv(ROOT/"tables"/"split_summary.csv")
tab(["Split","Wallets","Mean price","Median price","Price SD"],[[r['split'],f"{r['n']:,}",f"{r['target_mean']:.3f}",f"{r['target_median']:.3f}",f"{r['target_sd']:.3f}"] for r in split.to_dict("records")],[1.2,1.3,1.5,1.6,1.5])
para("We inspect development distributions using training wallets. The split summary documents that the final test prices are broadly comparable, and it does not guide model changes. No wallet receives a weight based on total spending, so a small number of large accounts cannot dominate the reported error.")
para("A lower average purchase price can reflect which markets and outcomes a wallet trades, or when it trades within those markets. Without market identities, resolutions, or trade timestamps, the target cannot identify bargain purchases, realized returns, or a preference for successful outcomes.")

section("Relationships among predictors")
para("Training rank correlations show strong overlap among several predictors. Total spending and largest purchase have a correlation of about 0.95. Trade count and market count have a correlation of about 0.79. Market count and concentration have a correlation of about minus 0.73. These relationships matter when interpreting coefficients and feature importance.")
fig("03_rank_correlations",width=6.45,max_height=5.45)
para("Individual monotonic correlations with average price are modest. Trade count has a rank correlation near 0.07 with the target, even though shuffling it later causes the largest increase in boosting error. Marginal correlation and predictive importance answer different questions. A weak one-variable association can still contribute to predictions alongside other inputs.")
para("Concentration also has a structural relationship with market count: it cannot fall below one divided by the number of markets. We therefore avoid treating a concentration and market-count correlation as a separate behavioral discovery.")

section("Observed price patterns across wallets")
para("We divide each training predictor into approximately equal-frequency bins and summarize average prices within each bin. The monetary predictors generally show higher observed prices in higher spending bins, while activity measures show curves that rise and then flatten. These plots motivate a flexible comparison but do not determine a causal effect or isolate a predictor from the others.")
fig("04_binned_associations",max_height=4.4)
para("The shaded ranges contain the middle half of observed wallet prices within each bin. Their width shows substantial variation among wallets with similar predictor values. They are not confidence intervals for the plotted means. Discrete counts can create repeated quantile boundaries, so some panels have fewer than ten bins.")
para("These summaries average over differences in all other wallet characteristics. They can therefore differ from conditional linear coefficients. For example, the marginal relation with trade count is slightly positive, while the fitted linear coefficient becomes negative after accounting for spending, market count, span, and concentration. That sign change should not be read as evidence that trading more causes lower prices.")

section("Training validation and the final comparison")
para("The boosting search compares six combinations of maximum leaves and L2 regularization on the same five training folds. Both the full and activity-only searches select 31 leaves with L2 regularization of one. The search is intentionally modest. The selected model is the best candidate within this grid rather than a globally optimized algorithm.")
fig("05_cross_validation",max_height=3.45)
tab(["Primary model","Mean validation MAE cents","Fold SD cents","Final test MAE cents"],[[n,f"{M[n]['CV_MAE']*100:.2f}",f"{M[n]['CV_MAE_SD']*100:.2f}",f"{M[n]['MAE']*100:.2f}"] for n in PRIMARY],[2.2,1.75,1.45,1.7])
para("Validation and test errors are similar, which supports the comparison within this cohort. Boosting training MAE is 16.04 cents, below its test MAE of 16.74 cents. This gap is consistent with some training optimism. We make no further model changes after reviewing the test diagnostics.")
tab(["MAE reduction from boosting","Estimate cents","95 percent paired interval"],[["Relative to mean baseline","5.41","5.15 to 5.66"],["Relative to linear regression","3.04","2.86 to 3.23"]],[3.15,1.45,2.5])
para("The paired comparison supports an improvement on these held out wallets under the bootstrap assumptions. The result does not show that a similar gain will hold for later periods, different events, or groups of dependent wallets.")

section("Predictions and residual patterns")
para("The observed-versus-predicted plots show where each model approximates the target and where predictions move toward the middle. Boosting spreads predictions more flexibly, but both models still make sizable individual errors. The diagonal marks perfect prediction. The displayed axes emphasize the target range.")
fig("07_observed_vs_predicted",max_height=3.35)
fig("08_residual_patterns",max_height=3.15)
para("Residuals equal observed minus predicted price. A positive residual means the model predicts too low. Because prices have lower and upper bounds, residual plots naturally have a constrained envelope. We report unclipped predictions and address the small number outside zero to one in the sensitivity section.")

section("The size and location of prediction errors")
para(f"Boosting predicts {D['boosting_error_within_5_cents']*100:.1f}% of test wallets within five cents, {D['boosting_error_within_10_cents']*100:.1f}% within ten cents, and {D['boosting_error_within_20_cents']*100:.1f}% within twenty cents. These figures put the average error in context. A lower error than the baseline still leaves substantial uncertainty for an individual wallet.")
fig("09_error_distributions",max_height=3.15)
fig("11_error_by_price",max_height=3.25)
para("Price-group boundaries come from training quintiles. Grouping the test wallets by their observed target helps describe where errors occur, but those target groups are unavailable before prediction. They cannot serve as prospective inputs. Average error also differs by price level, so one overall score does not describe every part of the distribution.")

section("What the models use")
para("We assess boosting reliance by shuffling one test predictor at a time while keeping the trained model fixed. The increase in MAE measures how much this fitted model depends on the information in that predictor under the shuffle procedure. We repeat each shuffle twenty times. The displayed variability reflects random shuffling, not a confidence interval for a causal effect. [R5]")
fig("12_permutation_importance",max_height=3.55)
para("Trade count produces the largest error increase at 6.32 cents. Total spending, largest purchase, and market count follow at roughly 3.4 to 3.8 cents. Observed span contributes less under this measure, while concentration has the smallest individual increase. Shuffling breaks relationships among correlated predictors, so the ranking is specific to this model and dataset.")
tab(["Linear model predictor","Price difference per training SD cents"],[[LABELS[r['variable']],f"{r['coefficient_cents']:+.2f}"] for r in S['coefficients']],[4.65,2.45])
para("Linear coefficients describe conditional associations after the stated transformations. They do not show effect sizes for an intervention. Shared information among predictors and the additive form limit their substantive interpretation. The companion coefficient figure provides the same comparison visually.")

section("Sensitivity to predictor choice and price bounds")
para("We prespecify a four-predictor sensitivity model using trade count, market count, observed span, and concentration. It removes total spending and largest purchase. Both boosting specifications receive the same six-candidate training-only search. This checks whether the main result relies heavily on monetary summaries computed from the same transactions as the target.")
fig("14_spending_ablation",max_height=3.65)
para("With only activity and concentration, boosting MAE rises from 16.74 to 19.97 cents and R squared falls from 0.368 to 0.148. Monetary inputs therefore add predictive information in this sample. Activity-only boosting still improves on the mean and median constants. The exercise remains a prediction of same-period summaries because every input comes from the observed history.")
sub("Alternative constant and bounded predictions")
para("Because MAE is the primary metric, we also check a training-median baseline, the optimal constant for absolute loss. Its test MAE is 21.55 cents. Boosting improves on it by 22.3%, so the main result does not depend on choosing the less favorable mean constant.")
tab(["Model","Predictions outside zero to one","MAE after clipping cents"],[[r['model'],r['clipped_predictions'],f"{r['MAE']*100:.3f}"] for r in S['clipping']],[2.3,2.8,2.0])
para("Only 22 linear predictions and 17 boosting predictions fall outside zero to one. Clipping changes the errors by less than one hundredth of a cent per share and preserves the ranking. We retain the original predictions for the primary comparison rather than choose a postprocessing rule using the test set.")

section("Variation across wallet groups and calibration")
para("We compare errors for one-market wallets, wallets with zero observed activity span, and wallets meeting the demo high-activity rule. These are descriptive, overlapping groups. No group is removed or used to tune a new model. The high-activity group contains only 50 test wallets, so its apparent performance has limited precision.")
fig("15_wallet_subgroups",max_height=3.7)
para("Boosting MAE is 20.78 cents among the 762 one-market test wallets, compared with 16.44 cents among the 10,429 multi-market wallets. For the 247 wallets with zero observed span, it is 19.15 cents. These comparisons show that the pooled result can conceal differences in predictability across wallet histories.")
sub("Average calibration")
para("The separate calibration plot groups test wallets into ten approximately equal-frequency bins based on predicted price. It compares each bin's mean prediction with its mean observed price. Its intervals use standard errors of observed bin means under wallet independence. These summaries diagnose average alignment. They do not assess prediction-interval coverage or guarantee small errors for individual wallets.")
para("The figure gallery contains the full calibration plot and its bin table. Diagnostic figures are used to describe the selected models. They do not trigger feature changes, additional tuning, or selection on the final test wallets.")

section("Interpretation and study limitations")
para("The six wallet summaries contain useful information about average purchase price. The flexible model provides a meaningful gain over an additive linear specification on the fixed test split, consistent with predictive value in nonlinear patterns or interactions. The experiment does not isolate which specific interaction produces that gain.")
sub("What the result describes")
para("The target and monetary inputs summarize the same underlying trades. This is a legitimate comparison for reconstructing a wallet-level summary when those inputs are already available. Its value as an economic forecast would require predictors observed before the future target period. The current table cannot support that temporal design.")
para("Wallets differ in which markets and outcomes they buy. Market mix and trading timing may explain some price variation, and neither is available in the aggregated table. Price alone cannot establish risk preference, prediction skill, profit, or access to private information. We make no claims about those outcomes.")
sub("Where uncertainty remains")
para("The sampled event groups and observation window constrain generalization. Wallets can share the same markets or respond to common news, and a person can control more than one wallet. Random wallet splitting does not test performance on entirely unseen markets or future events. The reported bootstrap treats wallets as independent and may understate uncertainty when errors are correlated.")
para("The MAE intervals condition on one fitted set of models and one split. They exclude the variation from rebuilding training samples or searching different parameter grids. The grid contains only six boosting configurations and compares one linear specification. Conclusions apply to this benchmark rather than every possible linear or nonlinear model.")
para("Predictor importance and coefficient signs depend on correlated inputs and model form. Permutation importance gives no direction of effect. The confirmed flag is excluded throughout, so the study neither detects insiders nor evaluates insider labels.")
sub("A defensible presentation takeaway")
para("Wallet activity and spending summaries predict part of the variation in average purchase price. Gradient boosting reduces held out MAE by 15.4% relative to linear regression, but a typical absolute error of about 16.7 cents leaves considerable individual uncertainty. The result describes associations within this selected historical cohort.")

section("Reproducibility and source record")
para("The accompanying analysis script loads only the agreed eight columns, verifies data constraints, saves wallet split membership, makes training-only development figures, fits models with training-only validation, and then saves fixed test predictions and diagnostics. Analysis outputs do not contain or use the ignored flag.")
tab(["Reproduction setting","Value"],[["Split and fold seed","42"],["Bootstrap and permutation seed","2026"],["Training folds","5 shuffled folds"],["Bootstrap resamples","2,000 paired wallet samples"],["Permutation repeats","20 per predictor"],["Python",S['versions']['python']],["scikit learn",S['versions']['scikit-learn']],["NumPy",S['versions']['numpy']],["pandas",S['versions']['pandas']],["Matplotlib",S['versions']['matplotlib']]],[3.4,3.7])
para("To reproduce the analysis, install code/requirements.txt in a Python 3.12 environment and run code/analyze_wallet_prices.py with --data pointing to the supplied CSV. The script writes results relative to its study folder. The source table remains unchanged. Saved fitted models are provided for reproducing predictions in the recorded package versions.")
para("The main numeric records are tables/primary_model_metrics.csv, tables/all_model_metrics.csv, tables/paired_bootstrap_comparisons.csv, tables/cv_grid_full.csv, tables/test_predictions.csv, and tables/analysis_summary.json. Wallet split membership is in tables/wallet_split.csv. PNG figures support direct reuse and SVG figures preserve editable vector elements.")
sub("Sources")
for key,title,source in REFERENCE:
    if source.startswith("https://"):
        p=DOC.add_paragraph(f"[{key}] ")
        link=OxmlElement("w:hyperlink")
        link.set(qn("r:id"),p.part.relate_to(source,RT.HYPERLINK,is_external=True))
        run=OxmlElement("w:r");rp=OxmlElement("w:rPr")
        color=OxmlElement("w:color");color.set(qn("w:val"),"315C8C");rp.append(color)
        size=OxmlElement("w:sz");size.set(qn("w:val"),"18");rp.append(size)
        run.append(rp);text=OxmlElement("w:t");text.text=title+" in the scikit learn 1.7 documentation";run.append(text)
        link.append(run);p._p.append(link)
        HTML.append(f'<p>[{key}] <a href="{html.escape(source)}">{html.escape(title)} in the scikit learn 1.7 documentation</a></p>')
        MD.append(f'[{key}] [{title}]({source})\n')
    else:
        p=para(f"[{key}] {title}. {source}")
    p.paragraph_format.space_after=Pt(5)
    for r in p.runs: r.font.size=Pt(9)

section("Figure index and presentation selection")
para("The complete figure gallery includes every figure below, with full captions. The report emphasizes the figures that explain the pipeline and main conclusions. The four-slide presentation selects the target distribution, model error comparison, and permutation importance. Its approach slide explains the split and comparison without adding an introduction or closing slide.")
tab(["Figure","Content","Use"],[[r['id'][:2],r['title'],"Presentation and report" if "slide" in r['audience'] else "Report or supplement"] for r in S['figures']],[.65,4.75,1.7])
para("All figure data and model diagnostics remain available in the tables folder. The presentation uses the same underlying values and model names as this report. Its editable charts preserve those values for reuse in the team presentation.")

DOC.save(OUT/"wallet_purchase_price_report.docx")
style="""body{font-family:Arial,sans-serif;color:#202c3d;line-height:1.6;max-width:1080px;margin:40px auto;padding:0 28px;background:#fff}h1{font-size:32px;color:#111}h2{font-size:26px;margin-top:54px;color:#111}h3{color:#111;margin-top:28px}p{max-width:94ch}table{border-collapse:collapse;width:100%;margin:24px 0;font-size:15px}th,td{padding:10px 12px;border:1px solid #d9d9d9;text-align:left}th{background:#263f5c;color:white}tr:nth-child(even){background:#f0f4f8}figure{margin:30px 0}img{width:100%;height:auto}figcaption{font-size:14px;color:#526073;margin-top:10px}a{color:#315c8c}code{background:#f0f4f8;padding:2px 4px}@media print{body{max-width:none;margin:0;padding:0}h2{break-before:page}figure,table{break-inside:avoid}}"""
(OUT/"wallet_purchase_price_report.html").write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>Wallet purchase price analysis</title><style>'+style+'</style><body>'+"\n".join(HTML)+"</body></html>")
(OUT/"wallet_purchase_price_report.md").write_text("\n".join(MD))
gallery=['<h1>Wallet purchase price figure gallery</h1><p>All 15 figures from the study. The confirmed flag is excluded throughout. Each figure is also saved separately as a PNG and an SVG in the figures folder.</p>']
for item in S['figures']:
    png=ROOT/"figures"/(item['id']+".png")
    b64=base64.b64encode(png.read_bytes()).decode()
    gallery.append('<h2>'+item['id'][:2]+" "+html.escape(item['title'])+'</h2><figure><img alt="'+html.escape(item['title'])+'" src="data:image/png;base64,'+b64+'"><figcaption>'+html.escape(item['caption'])+'</figcaption></figure>')
(OUT/"figure_gallery.html").write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>Wallet analysis figures</title><style>'+style+'</style><body>'+"\n".join(gallery)+"</body></html>")
print("Created report DOCX HTML and Markdown plus self contained figure gallery.")
